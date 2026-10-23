#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 prd.md 转成 docx（python-docx）。处理：标题/表格/无序列表/段落/加粗/代码/横线。"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BRAND = RGBColor(0x10, 0x63, 0x76)

def set_cn_font(run, name="微软雅黑"):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)

def add_md_to_doc(doc, md_text):
    lines = md_text.split("\n")
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
            i += 1
            continue

        # 空行 / 横线
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2  # 跳过分隔行
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for j, h in enumerate(header):
                hdr[j].text = ""
                run = hdr[j].paragraphs[0].add_run(h)
                run.bold = True
                run.font.size = Pt(9)
                set_cn_font(run)
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    cells[j].text = ""
                    run = cells[j].paragraphs[0].add_run(val)
                    run.font.size = Pt(9)
                    set_cn_font(run)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            set_cn_font(run)
            if level == 1:
                run.font.size = Pt(18)
                run.font.color.rgb = BRAND
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif level == 2:
                run.font.size = Pt(14)
                run.font.color.rgb = BRAND
            elif level == 3:
                run.font.size = Pt(12)
            else:
                run.font.size = Pt(11)
            i += 1
            continue

        # 无序列表
        m = re.match(r"^\s*[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\s*\d+[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

def add_inline(p, text):
    # 处理 **加粗** 和 `代码`
    pos = 0
    for m in re.finditer(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            set_cn_font(r)
        token = m.group(0)
        if token.startswith("**"):
            r = p.add_run(token[2:-2])
            r.bold = True
            set_cn_font(r)
        else:
            r = p.add_run(token[1:-1])
            r.font.name = "Menlo"
            r.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_cn_font(r)

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "prd.md"
    out = sys.argv[2] if len(sys.argv) > 2 else "prd.docx"
    md = open(src, encoding="utf-8").read()
    doc = Document()
    # 正文默认字号
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    add_md_to_doc(doc, md)
    doc.save(out)
    print("saved", out)

if __name__ == "__main__":
    main()
